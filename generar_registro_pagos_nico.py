#!/usr/bin/env python3
"""Genera Registro_Pagos_Nico.xlsx con openpyxl.

Planilla personal e independiente del presupuesto familiar: lleva la
contabilidad de todo lo que Nico va pagando del viaje (fecha, rubro, detalle,
monto), con el total por rubro y el total general calculados con formulas
nativas de Excel. Fuente de datos: pagos_nico.json, editado por Claude a
medida que Nico cuenta sus pagos por chat (no hay que cargar nada a mano).
"""
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT_PATH = "Registro_Pagos_Nico.xlsx"
PAGOS_PATH = "pagos_nico.json"


def cargar_pagos():
    if not os.path.exists(PAGOS_PATH):
        return []
    with open(PAGOS_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


# ---------------------------------------------------------------------------
# Estilos (mismos que Presupuesto_Florida_2027_Nico.xlsx, para consistencia)
# ---------------------------------------------------------------------------
NAVY = "1F3864"
BLUE = "2E5C8A"
LIGHT_BLUE = "D9E2F3"
GOLD = "FFC000"
GREY = "F2F2F2"
WHITE = "FFFFFF"

TITLE_FONT = Font(name="Calibri", size=18, bold=True, color=WHITE)
H2_FONT = Font(name="Calibri", size=11, bold=True, color=WHITE)
BODY_FONT = Font(name="Calibri", size=10.5)
BODY_BOLD = Font(name="Calibri", size=10.5, bold=True)

HEADER_FILL = PatternFill("solid", fgColor=NAVY)
SUBHEADER_FILL = PatternFill("solid", fgColor=BLUE)
LIGHT_FILL = PatternFill("solid", fgColor=LIGHT_BLUE)
TOTAL_FILL = PatternFill("solid", fgColor=GOLD)
STRIPE_FILL = PatternFill("solid", fgColor=GREY)

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

USD = '$#,##0.00'
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")

RUBROS = ["Vuelos", "Disney", "Universal", "Hospedaje Miami", "Auto/Transporte",
          "Comida", "Compras", "Seguro", "Propinas", "Otro"]

# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------
wb = Workbook()
ws = wb.active
ws.title = "Mis Pagos"
ws.sheet_view.showGridLines = False
widths = {"A": 13, "B": 17, "C": 46, "D": 15, "E": 34}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

ws.merge_cells("A1:E1")
c = ws.cell(row=1, column=1, value="REGISTRO DE PAGOS - NICOLAS")
c.font = TITLE_FONT
c.fill = HEADER_FILL
c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws.row_dimensions[1].height = 32
for col in range(1, 6):
    ws.cell(row=1, column=col).fill = HEADER_FILL

row = 2
ws.merge_cells(f"A{row}:E{row}")
c = ws.cell(row=row, column=1, value=(
    "Contabilidad personal de Nico: todo lo que va pagando del viaje, independiente de la Matriz "
    "Financiera familiar. Le contas cada pago a Claude por chat (fecha, rubro, detalle, monto) y esta "
    "hoja se actualiza sola. Los totales de abajo se recalculan solos con formulas."
))
c.font = Font(italic=True, size=10)
c.alignment = WRAP_TOP
c.fill = LIGHT_FILL
ws.row_dimensions[row].height = 34
row += 2

headers = ["Fecha", "Rubro", "Detalle del gasto", "Monto USD", "Notas"]
header_row = row
for j, h in enumerate(headers, start=1):
    cell = ws.cell(row=row, column=j, value=h)
    cell.font = H2_FONT
    cell.fill = SUBHEADER_FILL
    cell.alignment = CENTER
    cell.border = BOX
ws.row_dimensions[row].height = 18
ws.freeze_panes = f"A{row + 1}"
row += 1
pago_first_row = row

pagos = cargar_pagos()
N_BLANK_ROWS = 60
N_TOTAL_ROWS = max(len(pagos), 1) + N_BLANK_ROWS

for i in range(N_TOTAL_ROWS):
    r = row + i
    for col in range(1, 6):
        cell = ws.cell(row=r, column=col)
        cell.border = BOX
        cell.font = BODY_FONT
        if col == 4:
            cell.number_format = USD
        if (i % 2) == 0:
            cell.fill = STRIPE_FILL
    ws.row_dimensions[r].height = 16
    if i < len(pagos):
        p = pagos[i]
        ws.cell(row=r, column=1, value=p.get("fecha", ""))
        ws.cell(row=r, column=2, value=p.get("rubro", ""))
        ws.cell(row=r, column=3, value=p.get("detalle", ""))
        if p.get("monto") is not None:
            ws.cell(row=r, column=4, value=p["monto"])
        ws.cell(row=r, column=5, value=p.get("notas", ""))
        for col in (3, 5):
            ws.cell(row=r, column=col).alignment = WRAP_TOP
row += N_TOTAL_ROWS
pago_last_row = row - 1

rubro_list = '"' + ",".join(RUBROS) + '"'
dv_rubro = DataValidation(type="list", formula1=rubro_list, allow_blank=True, showDropDown=False)
ws.add_data_validation(dv_rubro)
dv_rubro.add(f"B{pago_first_row}:B{pago_last_row}")

row += 1
ws.merge_cells(f"A{row}:E{row}")
c = ws.cell(row=row, column=1, value="RESUMEN POR RUBRO (se calcula solo)")
c.font = H2_FONT
c.fill = SUBHEADER_FILL
ws.row_dimensions[row].height = 18
for col in range(1, 6):
    ws.cell(row=row, column=col).fill = SUBHEADER_FILL
row += 1

rango_rubro = f"$B${pago_first_row}:$B${pago_last_row}"
rango_monto = f"$D${pago_first_row}:$D${pago_last_row}"

for rubro in RUBROS:
    ws.cell(row=row, column=1, value=rubro).font = BODY_BOLD
    fcell = ws.cell(row=row, column=4, value=f'=SUMIF({rango_rubro},"{rubro}",{rango_monto})')
    fcell.number_format = USD
    fcell.font = BODY_BOLD
    for col in range(1, 6):
        ws.cell(row=row, column=col).border = BOX
    ws.row_dimensions[row].height = 16
    row += 1

ws.cell(row=row, column=1, value="TOTAL PAGADO POR NICO").font = Font(bold=True, size=12)
total_cell = ws.cell(row=row, column=4, value=f"=SUM({rango_monto})")
total_cell.number_format = USD
total_cell.font = Font(bold=True, size=12)
for col in range(1, 6):
    ws.cell(row=row, column=col).fill = TOTAL_FILL
    ws.cell(row=row, column=col).border = BOX
ws.row_dimensions[row].height = 22
row += 1

ws.cell(row=row, column=1, value="Cantidad de pagos registrados").font = Font(italic=True, size=9.5)
count_cell = ws.cell(row=row, column=4, value=f'=COUNTA({rango_rubro})')
count_cell.font = Font(italic=True, size=9.5)
row += 1

wb.save(OUTPUT_PATH)
print(f"OK: {OUTPUT_PATH} generado con {len(pagos)} pago(s) cargado(s).")
